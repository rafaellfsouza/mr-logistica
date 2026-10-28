"""
Exportação de relatórios para **Excel (.xlsx)**, **CSV** e **PDF**.

Todas as funções recebem a estrutura padrão gerada por :mod:`mrlog.relatorios` e
devolvem o caminho do arquivo criado. Nada depende de internet.
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from .utils import formatar_dinheiro, hoje, para_reais, slug

FORMATOS = ("xlsx", "csv", "pdf")


class ErroExportacao(Exception):
    """Falha ao exportar (formato inválido, sem permissão de escrita etc.)."""


def _preparar(dados: dict, formato: str) -> tuple[Path, list[str], list[list[Any]]]:
    caminho = Path(dados.get("arquivo") or "")
    if not caminho.name:
        raise ErroExportacao("Informe o caminho do arquivo de destino.")
    if caminho.suffix.lower().lstrip(".") != formato:
        caminho = caminho.with_suffix("." + formato)
    caminho.parent.mkdir(parents=True, exist_ok=True)

    cabecalho = [str(c["rotulo"]) for c in dados["colunas"]]
    linhas: list[list[Any]] = []
    for linha in dados["linhas"]:
        valores = []
        for coluna in dados["colunas"]:
            v = linha.get(coluna["chave"])
            tipo = coluna["tipo"]
            if v is None:
                valores.append("")
            elif tipo == "dinheiro":
                valores.append(para_reais(v))
            elif tipo == "quantidade":
                valores.append(float(v))
            elif tipo == "inteiro":
                valores.append(int(v))
            elif tipo == "percentual":
                valores.append(float(v) / 100.0)
            elif tipo == "data":
                from .utils import formatar_data

                valores.append(formatar_data(v))
            else:
                valores.append(str(v))
        linhas.append(valores)
    return caminho, cabecalho, linhas


# --------------------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------------------


def exportar_excel(dados: dict, arquivo) -> str:
    """Gera um arquivo ``.xlsx`` formatado (título, cabeçalho, filtros e totais)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - dependência obrigatória no build
        raise ErroExportacao("Biblioteca openpyxl não instalada. Reinstale o aplicativo.") from exc

    destino = Path(arquivo)
    if destino.suffix.lower() != ".xlsx":
        destino = destino.with_suffix(".xlsx")
    destino.parent.mkdir(parents=True, exist_ok=True)

    cabecalho, linhas = [], []
    caminho, cabecalho, linhas = _preparar({**dados, "arquivo": destino}, "xlsx")

    livro = Workbook()
    planilha = livro.active
    planilha.title = (dados.get("chave") or "Relatório")[:31]

    fino = Side(style="thin", color="D0D0D0")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    azul = "1F4E79"

    # Título e período
    planilha["A1"] = dados.get("titulo", "Relatório")
    planilha["A1"].font = Font(bold=True, size=14, color=azul)
    planilha["A2"] = f"Período: {dados.get('periodo', '')}    |    Gerado em: {hoje()}"
    planilha["A2"].font = Font(italic=True, size=9, color="666666")
    if dados.get("filtros"):
        planilha["A3"] = f"Filtros: {dados['filtros']}"
        planilha["A3"].font = Font(size=9, color="666666")
        linha_cabecalho = 5
    else:
        linha_cabecalho = 4

    for i, titulo in enumerate(cabecalho, start=1):
        celula = planilha.cell(row=linha_cabecalho, column=i, value=titulo)
        celula.font = Font(bold=True, color="FFFFFF", size=10)
        celula.fill = PatternFill("solid", fgColor=azul)
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celula.border = borda

    formatos = {c["chave"]: c["tipo"] for c in dados["colunas"]}
    for r, linha in enumerate(linhas, start=linha_cabecalho + 1):
        for c, valor in enumerate(linha, start=1):
            celula = planilha.cell(row=r, column=c, value=valor)
            celula.border = borda
            tipo = dados["colunas"][c - 1]["tipo"]
            if tipo == "dinheiro":
                celula.number_format = 'R$ #,##0.00'
                celula.alignment = Alignment(horizontal="right")
            elif tipo == "quantidade":
                celula.number_format = "#,##0.0000"
                celula.alignment = Alignment(horizontal="right")
            elif tipo == "percentual":
                celula.number_format = "0.00%"
                celula.alignment = Alignment(horizontal="right")
            elif tipo == "inteiro":
                celula.number_format = "#,##0"
                celula.alignment = Alignment(horizontal="right")
            else:
                celula.alignment = Alignment(horizontal="left", vertical="center")
        if (r - linha_cabecalho) % 2 == 0:
            for c in range(1, len(cabecalho) + 1):
                planilha.cell(row=r, column=c).fill = PatternFill("solid", fgColor="F2F6FA")

    ultima_linha = linha_cabecalho + len(linhas)

    # Linha de totais
    totais = dados.get("totais") or {}
    if totais:
        r = ultima_linha + 1
        planilha.cell(row=r, column=1, value="TOTAIS").font = Font(bold=True)
        coluna = 2
        for chave, valor in totais.items():
            celula = planilha.cell(row=r, column=coluna, value=para_reais(valor) if isinstance(valor, int) and abs(valor) > 99 else valor)
            celula.font = Font(bold=True)
            if isinstance(valor, int) and abs(valor) > 99:
                celula.number_format = 'R$ #,##0.00'
            coluna += 1
        for c in range(1, len(cabecalho) + 1):
            planilha.cell(row=r, column=c).fill = PatternFill("solid", fgColor="DCE6F1")
            planilha.cell(row=r, column=c).border = borda
        ultima_linha = r

    # Resumo em texto
    if dados.get("resumo"):
        r = ultima_linha + 2
        for texto in dados["resumo"]:
            planilha.cell(row=r, column=1, value=texto).font = Font(size=9, italic=True, color="444444")
            r += 1

    # Larguras, filtros e painel congelado
    for i, coluna in enumerate(dados["colunas"], start=1):
        letra = get_column_letter(i)
        planilha.column_dimensions[letra].width = min(max(int(coluna.get("largura") or 12), 9), 46)
    if len(linhas) > 0:
        planilha.auto_filter.ref = (
            f"A{linha_cabecalho}:{get_column_letter(len(cabecalho))}{linha_cabecalho + len(linhas)}"
        )
    planilha.freeze_panes = planilha.cell(row=linha_cabecalho + 1, column=1)
    planilha.sheet_view.showGridLines = False

    try:
        livro.save(str(caminho))
    except OSError as exc:
        raise ErroExportacao(f"Não foi possível salvar o arquivo: {exc}") from exc
    return str(caminho)


# --------------------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------------------


def exportar_csv(dados: dict, arquivo, separador: str = ";") -> str:
    """Gera um CSV compatível com Excel em português (``;`` e vírgula decimal)."""
    caminho, cabecalho, linhas = _preparar({**dados, "arquivo": arquivo}, "csv")
    try:
        with open(caminho, "w", newline="", encoding="utf-8-sig") as arq:
            escritor = csv.writer(arq, delimiter=separador, quoting=csv.QUOTE_MINIMAL)
            escritor.writerow([dados.get("titulo", "Relatório")])
            escritor.writerow([f"Período: {dados.get('periodo', '')}"])
            escritor.writerow([f"Gerado em: {hoje()}"])
            escritor.writerow([])
            escritor.writerow(cabecalho)
            for linha in linhas:
                escritor.writerow(
                    [
                        (str(v).replace(".", ",") if isinstance(v, float) else v)
                        for v in linha
                    ]
                )
            if dados.get("totais"):
                escritor.writerow([])
                escritor.writerow(["TOTAIS"] + [
                    f"{k}: {formatar_dinheiro(v) if isinstance(v, int) and abs(v) > 99 else v}"
                    for k, v in dados["totais"].items()
                ])
    except OSError as exc:
        raise ErroExportacao(f"Não foi possível salvar o CSV: {exc}") from exc
    return str(caminho)


# --------------------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------------------


def _escapar(texto: Any) -> str:
    """Escapa os caracteres especiais do Paragraph (reportlab usa mini-XML)."""
    return (
        str(texto)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def exportar_pdf(dados: dict, arquivo, paisagem: bool | None = None, titulo_extra: str = "") -> str:
    """Gera o relatório em PDF (A4, com cabeçalho, tabela e rodapé numerado)."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            HRFlowable,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as exc:  # pragma: no cover
        raise ErroExportacao("Biblioteca reportlab não instalada. Reinstale o aplicativo.") from exc

    destino = Path(arquivo)
    if destino.suffix.lower() != ".pdf":
        destino = destino.with_suffix(".pdf")
    destino.parent.mkdir(parents=True, exist_ok=True)
    _, cabecalho, linhas = _preparar({**dados, "arquivo": destino}, "pdf")
    colunas = dados["colunas"]

    if paisagem is None:
        paisagem = len(colunas) > 7
    pagina = landscape(A4) if paisagem else A4
    margem = 12 * mm
    largura_util = pagina[0] - 2 * margem

    estilos = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle("t", parent=estilos["Title"], fontSize=14, leading=17, spaceAfter=2, textColor=colors.HexColor("#1F4E79"))
    estilo_sub = ParagraphStyle("s", parent=estilos["Normal"], fontSize=8.5, leading=11, textColor=colors.HexColor("#555555"))
    estilo_celula = ParagraphStyle("c", parent=estilos["Normal"], fontSize=7.4, leading=9)
    estilo_celula_dir = ParagraphStyle("cd", parent=estilo_celula, alignment=TA_RIGHT)
    estilo_cabecalho = ParagraphStyle("h", parent=estilo_celula, textColor=colors.white, fontSize=7.6, leading=9.2)

    # Distribui a largura disponível proporcionalmente ao tamanho declarado de cada coluna
    pesos = [max(6, int(c.get("largura") or 10)) for c in colunas]
    total_pesos = sum(pesos)
    larguras = [largura_util * p / total_pesos for p in pesos]

    def formatar_celula(valor: Any, tipo: str, alinhamento: str) -> Paragraph:
        if tipo == "dinheiro":
            texto = formatar_dinheiro(int(valor or 0)) if isinstance(valor, (int, float)) else str(valor)
        elif tipo == "percentual":
            texto = f"{float(valor or 0) * 100:.2f}%".replace(".", ",") if isinstance(valor, (int, float)) else str(valor)
        elif tipo == "quantidade":
            from .utils import formatar_quantidade

            texto = formatar_quantidade(valor) if isinstance(valor, (int, float)) else str(valor)
        elif tipo == "inteiro":
            texto = f"{int(valor or 0)}" if isinstance(valor, (int, float)) else str(valor)
        else:
            texto = "" if valor is None else str(valor)
        estilo = estilo_celula_dir if (alinhamento == "direita" or tipo in ("dinheiro", "quantidade", "percentual", "inteiro")) else estilo_celula
        return Paragraph(_escapar(texto), estilo)

    tabela_dados = [[Paragraph(_escapar(c), estilo_cabecalho) for c in cabecalho]]
    for linha in linhas:
        tabela_dados.append(
            [formatar_celula(v, colunas[i]["tipo"], colunas[i].get("alinhamento", "")) for i, v in enumerate(linha)]
        )

    totais = dados.get("totais") or {}
    if totais:
        linha_total = [Paragraph("<b>TOTAIS</b>", estilo_celula)] + [Paragraph("", estilo_celula)] * (len(cabecalho) - 1)
        indice = 1
        for chave, valor in totais.items():
            if indice >= len(linha_total):
                break
            texto = formatar_dinheiro(valor) if isinstance(valor, int) and abs(valor) > 99 else str(valor)
            linha_total[indice] = Paragraph(f"<b>{_escapar(texto)}</b>", estilo_celula_dir)
            indice += 1
        tabela_dados.append(linha_total)

    estilo_tabela = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9D3E0")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2.4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.4),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]
    )
    for i in range(1, len(tabela_dados)):
        if i % 2 == 0:
            estilo_tabela.add("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F2F6FA"))
    if totais:
        estilo_tabela.add("BACKGROUND", (0, len(tabela_dados) - 1), (-1, len(tabela_dados) - 1), colors.HexColor("#DCE6F1"))

    historia = [
        Paragraph(_escapar(dados.get("titulo", "Relatório")), estilo_titulo),
        Paragraph(
            _escapar(
                f"Período: {dados.get('periodo', '') or '—'}"
                + (f" &nbsp;|&nbsp; {titulo_extra}" if titulo_extra else "")
                + f" &nbsp;|&nbsp; Gerado em {hoje()}"
            ),
            estilo_sub,
        ),
    ]
    if dados.get("filtros"):
        historia.append(Paragraph(_escapar(f"Filtros: {dados['filtros']}"), estilo_sub))
    historia += [Spacer(1, 4), HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#1F4E79")), Spacer(1, 6)]
    historia.append(
        Table(tabela_dados, colWidths=larguras, repeatRows=1, style=estilo_tabela, hAlign=TA_LEFT)
    )

    if dados.get("resumo"):
        historia += [Spacer(1, 10), HRFlowable(width="100%", thickness=0.4, color=colors.HexColor("#BBBBBB")), Spacer(1, 4)]
        for texto in dados["resumo"]:
            historia.append(Paragraph(_escapar(texto), estilo_sub))
    if not linhas:
        historia.append(Spacer(1, 10))
        historia.append(Paragraph("<i>Nenhum registro encontrado para o período/filtro selecionado.</i>", estilo_sub))

    def rodape(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(margem, 8 * mm, "MR Logística — relatório gerado localmente (sem envio de dados)")
        canvas.drawRightString(pagina[0] - margem, 8 * mm, f"Página {doc.page}")
        canvas.restoreState()

    try:
        documento = SimpleDocTemplate(
            str(destino), pagesize=pagina, leftMargin=margem, rightMargin=margem,
            topMargin=margem, bottomMargin=14 * mm,
            title=dados.get("titulo", "Relatório"), author="MR Logística",
        )
        documento.build(historia, onFirstPage=rodape, onLaterPages=rodape)
    except Exception as exc:  # reportlab pode falhar com dados muito largos
        raise ErroExportacao(f"Falha ao gerar o PDF: {exc}") from exc
    return str(destino)


# --------------------------------------------------------------------------------------
# Despachante
# --------------------------------------------------------------------------------------


def exportar(dados: dict, arquivo, formato: str = "xlsx") -> str:
    """Exporta o relatório para ``xlsx``, ``csv`` ou ``pdf`` e devolve o caminho."""
    fmt = str(formato or "xlsx").lower().lstrip(".")
    if fmt in ("xls", "excel", "xlsx"):
        return exportar_excel(dados, arquivo)
    if fmt == "csv":
        return exportar_csv(dados, arquivo)
    if fmt == "pdf":
        return exportar_pdf(dados, arquivo)
    raise ErroExportacao(f"Formato de exportação não suportado: '{formato}'. Use xlsx, csv ou pdf.")


def caminho_sugerido(titulo: str, formato: str = "xlsx", pasta=None) -> Path:
    """Sugere um nome de arquivo: ``vendas-por-dia-2026-10-28.xlsx``."""
    from .caminhos import pasta_temporaria

    base = Path(pasta) if pasta else pasta_temporaria()
    base.mkdir(parents=True, exist_ok=True)
    return base / f"{slug(titulo)}-{hoje()}.{formato}"
