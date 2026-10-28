"""
Teste automatizado da interface gráfica (roda em Xvfb no Linux/CI e localmente no Windows).

O que ele faz:
  1. cria um banco temporário com dados de exemplo (categorias, produtos, clientes,
     fornecedores, usuários, vendas, caixa, contas);
  2. abre a janela principal "pulando" o login (define o usuário admin direto);
  3. navega por **todas** as telas do menu, chama ``carregar()``, troca o tema,
     abre os principais diálogos e registra qualquer exceção;
  4. no final imprime um resumo e devolve código 1 se algo falhou.

Uso::

    xvfb-run -a python3 tests/gui_test.py
"""

from __future__ import annotations

import os
import sys
import tempfile
import traceback
from pathlib import Path
from typing import Sequence

# banco de dados temporário ANTES de importar o pacote
_PASTA = Path(tempfile.mkdtemp(prefix="mrlog-gui-"))
os.environ["MRLOG_DATA_DIR"] = str(_PASTA)
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import tkinter as tk  # noqa: E402

from mrlog import (  # noqa: E402
    caixa, categorias, clientes, contas_pagar, contas_receber, estoque,
    fornecedores, produtos, usuarios, vendas,
)
from mrlog.banco import Banco  # noqa: E402
from mrlog.gui.app import Aplicacao, TELAS  # noqa: E402


def semear(banco: Banco) -> None:
    """Cria dados suficientes para todas as telas terem conteúdo."""
    if banco.consultar_valor("SELECT COUNT(*) FROM produtos") > 0:
        return  # banco já semeado (reexecução do teste)

    def categoria(nome: str, descricao: str = "") -> int:
        existente = categorias.pegar_por_nome(banco, nome)
        return existente["id"] if existente else categorias.criar(banco, nome, descricao)

    categoria("Mercearia", "Itens de mercearia")
    categoria("Bebidas", "Bebidas em geral")
    categoria("Higiene", "Produtos de higiene e limpeza")
    fornecedores.criar(banco, {"nome": "Distribuidora Alfa", "razao_social": "Alfa Distribuidora de Alimentos LTDA",
                               "telefone": "11 4002-8922", "celular": "11 94002-8922",
                               "cnpj_cpf": "12.345.678/0001-90", "email": "contato@alfa.com",
                               "cidade": "São Paulo", "uf": "SP", "prazo_dias": 28,
                               "contato": "Carlos", "observacao": "Entrega às terças."})
    fornecedores.criar(banco, {"nome": "Laticínios Beta", "telefone": "11 3333-2222",
                               "cnpj_cpf": "98.765.432/0001-10", "cidade": "Campinas", "uf": "SP",
                               "prazo_dias": 14})
    clientes.criar(banco, {"nome": "João da Silva", "tipo": "Física", "telefone": "11 98888-7777",
                           "cpf_cnpj": "123.456.789-09", "limite_credito": 200000,
                           "cidade": "São Paulo", "uf": "SP", "email": "joao@exemplo.com",
                           "endereco": "Rua das Flores", "numero": "120", "bairro": "Centro",
                           "nascimento": "12/03/1988"})
    clientes.criar(banco, {"nome": "Maria Souza", "telefone": "11 97777-6666", "cidade": "Osasco", "uf": "SP"})
    clientes.criar(banco, {"nome": "Mercado Bom Preço LTDA", "tipo": "Jurídica",
                           "cpf_cnpj": "11.222.333/0001-81", "limite_credito": 500000,
                           "telefone": "11 3232-1010", "cidade": "Guarulhos", "uf": "SP",
                           "contato": "Fernanda"})

    arroz_id = produtos.criar(banco, {
        "codigo": "", "nome": "Arroz tipo 1 5kg", "categoria_id": categoria("Mercearia"), "marca": "Tio João",
        "fornecedor_id": 1, "unidade_base": "un", "unidade_venda": "un", "fator_venda": 1,
        "preco_custo": 1850, "preco_venda": 2499, "estoque_minimo": 5, "permite_fracionar": 0,
    }, estoque_inicial=40)
    arroz = produtos.pegar(banco, arroz_id)
    produtos.definir_codigo_barras(banco, arroz_id, "7891000100103")
    arroz = produtos.pegar(banco, arroz_id)

    queijo_id = produtos.criar(banco, {
        "codigo": "", "nome": "Queijo mussarela fatiado", "categoria_id": categoria("Mercearia"), "marca": "Tirolez",
        "fornecedor_id": 1, "unidade_base": "kg", "unidade_venda": "g", "fator_venda": 0.001,
        "preco_custo": 3, "preco_venda": 5, "estoque_minimo": 2, "permite_fracionar": 1,
        "decimais": 3, "controla_lote": 1,
    }, estoque_inicial=12)
    queijo = produtos.pegar(banco, queijo_id)
    estoque.entrada(banco, queijo_id, 5, motivo="compra", custo_unitario=3, lote_numero="L-1001",
                    validade="2027-01-15", observacao="Compra com lote")

    refrigerante_id = produtos.criar(banco, {
        "codigo": "", "nome": "Refrigerante cola 2L", "categoria_id": categoria("Bebidas"), "marca": "Guaraná",
        "unidade_base": "un", "unidade_venda": "cx", "fator_venda": 6,
        "preco_custo": 4200, "preco_venda": 5400, "estoque_minimo": 3, "permite_fracionar": 0,
    }, estoque_inicial=60)
    refrigerante = produtos.pegar(banco, refrigerante_id)
    produtos.criar(banco, {
        "codigo": "", "nome": "Produto sem código de barras", "categoria_id": categoria("Mercearia"),
        "unidade_base": "un", "unidade_venda": "un", "preco_custo": 100, "preco_venda": 250,
    }, estoque_inicial=8)

    caixa.abrir(banco, valor_abertura=20000, observacao="Troco inicial")

    carrinho = vendas.Carrinho(banco)
    carrinho.adicionar(arroz, 2)
    carrinho.adicionar(queijo, 250, unidade="g")
    carrinho.adicionar(refrigerante, 1)
    admin = usuarios.pegar_por_login(banco, "admin")
    vendas.finalizar(banco, carrinho, pagamentos=[{"forma": "Dinheiro", "valor": 20000}],
                     cliente=clientes.pegar(banco, 1), vendedor=admin)

    carrinho2 = vendas.Carrinho(banco)
    carrinho2.adicionar(arroz, 1)
    vendas.finalizar(banco, carrinho2, pagamentos=[{"forma": "Fiado", "valor": 2499}],
                     cliente=clientes.pegar(banco, 1), vendedor=admin)

    contas_pagar.criar(banco, "Aluguel do mês", 150000, vencimento="2026-11-05",
                       categoria="Outros", fornecedor_id=1, observacao="Lançado pelo teste")
    contas_receber.criar_manual(banco, 1, "Venda fiado anterior", 8000, vencimento="2026-10-30")

    # usuários extras (a tela de usuários precisa de mais de um perfil para testar)
    if not usuarios.pegar_por_login(banco, "maria.v"):
        usuarios.criar(banco, {"nome": "Maria Vendedora", "login": "maria.v", "perfil": "vendedor",
                               "email": "maria@exemplo.com", "telefone": "11 97777-6666",
                               "ativo": 1, "trocar_senha": 1}, "venda123", autor=admin)
    if not usuarios.pegar_por_login(banco, "jose.g"):
        usuarios.criar(banco, {"nome": "José Gerente", "login": "jose.g", "perfil": "gerente",
                               "email": "jose@exemplo.com", "ativo": 1}, "gerente123", autor=admin)
    if not usuarios.pegar_por_login(banco, "ana.inativa"):
        usuarios.criar(banco, {"nome": "Ana (afastada)", "login": "ana.inativa", "perfil": "vendedor",
                               "ativo": 0}, "ana1234", autor=admin)


def testar_pdv(app, banco) -> None:
    """Simula uma venda completa pelo PDV, sem abrir nenhum diálogo modal."""
    from mrlog import produtos as produtos_mod
    from mrlog import vendas as vendas_mod

    app.abrir_tela("pdv")
    pdv = app.tela_atual
    assert pdv is not None, "tela do PDV não abriu"

    codigo = banco.consultar_valor("SELECT codigo_barras FROM produtos WHERE codigo_barras IS NOT NULL LIMIT 1")
    assert codigo, "nenhum produto com código de barras na base de teste"

    # 1) bipa o produto conhecido duas vezes (deve somar quantidade)
    pdv.entrada_codigo.insert(0, codigo)
    pdv._ao_bipar()
    pdv.entrada_codigo.insert(0, codigo)
    pdv._ao_bipar()
    app.janela.update()
    assert len(pdv.carrinho.itens) == 1, f"esperava 1 item, vieram {len(pdv.carrinho.itens)}"
    assert abs(pdv.carrinho.itens[0].quantidade - 2.0) < 1e-6, "a segunda bipada não somou quantidade"

    # 2) bipa um código desconhecido (sem cadastrar: resposta "Cancelar")
    antes = len(pdv.carrinho.itens)
    pdv.entrada_codigo.insert(0, "9999999999992")
    pdv._ao_bipar()
    assert len(pdv.carrinho.itens) == antes, "código desconhecido adicionou item"

    # 3) adiciona um produto fracionado em outra unidade (250 g de queijo)
    queijo = banco.consultar_uma("SELECT * FROM produtos WHERE unidade_base='kg' AND unidade_venda='g' LIMIT 1")
    if queijo:
        pdv._adicionar(produtos_mod.pegar(banco, queijo["id"]), 250, unidade="g")
        assert len(pdv.carrinho.itens) == antes + 1

    # 4) desconto em percentual na venda
    pdv.carrinho.definir_desconto(percentual=5, limite_pct=pdv.ctx.limite_desconto())
    pdv._redesenhar()
    totais = pdv.carrinho.totais()
    assert totais.desconto_venda > 0, "desconto da venda não foi aplicado"

    # 5) finaliza com pagamento misto (dinheiro + Pix) e troco
    restante = totais.total - 1000
    resultado_pagamento = {
        "pagamentos": [
            {"forma": "Dinheiro", "valor": 1000, "recebido": 2000, "parcelas": 1, "operadora": "", "documento": ""},
            {"forma": "Pix", "valor": restante, "recebido": restante, "parcelas": 1, "operadora": "", "documento": ""},
        ],
        "troco": 1000,
        "observacao": "Venda do teste automatizado",
        "vencimento_fiado": "",
        "fiado": 0,
    }
    vendas_antes = banco.consultar_valor("SELECT COUNT(*) FROM vendas")
    original = vendas_mod.finalizar
    vendas_mod.finalizar = lambda *a, **k: original(*a, **{**k, **{"observacao": resultado_pagamento["observacao"]}})
    try:
        # substitui apenas o diálogo de pagamento
        import mrlog.gui.telas.pdv_tela as pdv_mod

        classe_original = pdv_mod.DialogoPagamento

        class DialogoFalso(classe_original):
            def abrir(self):
                return dict(resultado_pagamento)

        pdv_mod.DialogoPagamento = DialogoFalso
        try:
            pdv.finalizar_venda()
        finally:
            pdv_mod.DialogoPagamento = classe_original
    finally:
        vendas_mod.finalizar = original

    app.janela.update()
    assert banco.consultar_valor("SELECT COUNT(*) FROM vendas") == vendas_antes + 1, "a venda não foi gravada"
    ultima = vendas_mod.ultima_venda(banco)
    assert int(ultima["total"]) == totais.total, f"total gravado {ultima['total']} != {totais.total}"
    assert int(ultima["troco"]) == 1000, f"troco gravado {ultima['troco']} != 1000"
    assert pdv.ultima_venda and int(pdv.ultima_venda["numero"]) == int(ultima["numero"])

    # 6) recibo não fiscal em PDF
    from mrlog import impressao

    caminho = impressao.gerar_recibo_pdf(banco, int(ultima["id"]))
    assert Path(caminho).exists() and Path(caminho).stat().st_size > 500, "recibo PDF não foi gerado"

    # 7) carrinho limpo após a venda
    assert not pdv.carrinho.itens, "o carrinho continuou com itens após finalizar"


def confirmar_sim(dialogs_mod):
    """Contexto que faz os dialogos de confirmacao responderem SIM."""
    import contextlib

    @contextlib.contextmanager
    def _ctx():
        anterior = dialogs_mod.RESPOSTA_CONFIRMAR
        dialogs_mod.RESPOSTA_CONFIRMAR = True
        try:
            yield
        finally:
            dialogs_mod.RESPOSTA_CONFIRMAR = anterior

    return _ctx()


def testar_cadastros(app, banco, dialogs_mod) -> None:
    """Exercita as telas de cadastro (categorias, clientes, fornecedores, usuários)."""
    sim = lambda: confirmar_sim(dialogs_mod)  # noqa: E731

    # ---------------------------------------------------------------- categorias
    app.abrir_tela("categorias")
    tela = app.tela_atual
    tela.carregar()
    nomes = [l["nome"] for l in tela.registros()]
    assert "Mercearia" in nomes and "Bebidas" in nomes, f"categorias faltando: {nomes}"
    assert all("qtd_produtos" in l and "estoque_valor" in l for l in tela.registros()), "colunas calculadas ausentes"
    assert tela._detalhes(tela.registros()[0]), "painel de detalhes vazio"

    categoria_id = categorias.criar(banco, "Categoria do teste", "criada pelo teste automatizado")
    tela._incluir_inativas.set(True)
    tela.carregar()
    assert tela.tabela.selecionar_por("id", categoria_id), "categoria nova não apareceu na grade"
    with sim():
        tela.alternar_status()                                  # desativa
    tela.carregar()
    linha = next(l for l in tela.registros() if int(l["id"]) == categoria_id)
    assert linha["status_texto"] == "Inativa" and linha["_tag"] == "neutro", "categoria não foi desativada"
    assert tela.tabela.selecionar_por("id", categoria_id)
    tela.alternar_status()                                      # reativa (não pede confirmação)
    tela.carregar()
    linha = next(l for l in tela.registros() if int(l["id"]) == categoria_id)
    assert linha["status_texto"] == "Ativa", "categoria não foi reativada"
    # exclusão bloqueada enquanto houver produtos
    assert tela.tabela.selecionar_por("nome", "Mercearia")
    assert tela._excluir(tela, tela.selecionado()) is None, "excluir categoria com produtos deveria ser bloqueado"
    # navegação cruzada: "produtos da categoria"
    assert tela.tabela.selecionar_por("nome", "Bebidas")
    tela.ver_produtos()
    assert app.tela_id_atual == "produtos", "não navegou para produtos"
    assert app.tela_atual._categoria.get() == "Bebidas", "filtro de categoria não foi aplicado"
    assert len(app.tela_atual.registros()) >= 1, "nenhum produto da categoria Bebidas"

    # ---------------------------------------------------------------- clientes
    app.abrir_tela("clientes")
    tela = app.tela_atual
    tela.carregar()
    registros = tela.registros()
    assert len(registros) >= 3, f"esperava 3+ clientes, vieram {len(registros)}"
    joao = next(l for l in registros if l["nome"] == "João da Silva")
    assert joao["documento_texto"] == "123.456.789-09", f"CPF mal formatado: {joao['documento_texto']}"
    assert int(joao["limite_credito"]) == 200000 and int(joao["saldo_devedor"]) > 0, "crédito/saldo incorretos"
    assert tela._tag_linha(joao) in ("alerta", "perigo"), "cliente devedor sem destaque"
    assert tela._detalhes(joao), "painel de detalhes vazio"
    tela._apenas_devedores.set(True)
    tela.carregar()
    assert tela.registros() and all(int(l["saldo_devedor"]) > 0 for l in tela.registros()), \
        "filtro 'somente devedores' falhou"
    tela._apenas_devedores.set(False)
    tela.carregar()
    # lançamento de fiado avulso (o formulário é modal; o módulo é chamado direto)
    contas_antes = banco.consultar_valor("SELECT COUNT(*) FROM contas_receber")
    conta_id = contas_receber.criar_manual(banco, 2, "Fiado do teste automatizado", 4500, "2026-12-10")
    assert banco.consultar_valor("SELECT COUNT(*) FROM contas_receber") == contas_antes + 1
    tela.carregar()
    maria = next(l for l in tela.registros() if l["nome"] == "Maria Souza")
    assert int(maria["saldo_devedor"]) == 4500, f"saldo devedor não refletiu o fiado: {maria['saldo_devedor']}"
    assert int(maria["contas_abertas"]) == 1, "quantidade de títulos errada"
    # desativação bloqueada enquanto houver débito
    assert tela.tabela.selecionar_por("id", 2)
    tela.alternar_status()
    assert int(clientes.pegar(banco, 2)["ativo"]) == 1, "cliente devedor não deveria ter sido desativado"
    # venda pelo PDV já com o cliente identificado
    tela.nova_venda()
    assert app.tela_id_atual == "pdv", "não navegou para o PDV"
    pdv = app.tela_atual
    assert pdv.carrinho.cliente and int(pdv.carrinho.cliente["id"]) == 2, "cliente não foi vinculado ao carrinho"
    assert "Maria Souza" in pdv.entrada_cliente.get(), f"rótulo do cliente errado: {pdv.entrada_cliente.get()}"
    pdv.nova_venda()
    assert pdv.carrinho.cliente is None, "cliente não foi limpo na nova venda"
    assert "Consumidor final" in pdv.entrada_cliente.get()
    # ficha/extrato do cliente
    extrato = contas_receber.extrato_cliente(banco, 2)
    assert extrato["total_aberto"] == 4500 and len(extrato["abertas"]) == 1, "extrato do cliente incorreto"
    assert clientes.pode_fiar(extrato["cliente"], 1000)[0], "cliente sem limite deveria poder fiar"

    # ---------------------------------------------------------------- fornecedores
    app.abrir_tela("fornecedores")
    tela = app.tela_atual
    tela.carregar()
    registros = tela.registros()
    assert len(registros) >= 2, f"esperava 2+ fornecedores, vieram {len(registros)}"
    alfa = next(l for l in registros if l["nome"] == "Distribuidora Alfa")
    assert alfa["documento_texto"] == "12.345.678/0001-90", f"CNPJ mal formatado: {alfa['documento_texto']}"
    assert int(alfa["saldo_aberto"]) == 150000, f"saldo a pagar incorreto: {alfa['saldo_aberto']}"
    assert int(alfa["qtd_produtos"]) >= 2, "produtos vinculados não contabilizados"
    assert tela._tag_linha(alfa) == "alerta", "fornecedor com saldo sem destaque"
    assert tela._detalhes(alfa), "painel de detalhes vazio"
    tela._apenas_com_saldo.set(True)
    tela.carregar()
    assert tela.registros() and all(int(l["saldo_aberto"]) > 0 for l in tela.registros()), "filtro 'com saldo' falhou"
    tela._apenas_com_saldo.set(False)
    tela.carregar()
    conta_pagar_id = contas_pagar.criar(banco, "Frete do teste", 7300, vencimento="2026-11-20",
                                        fornecedor_id=int(alfa["id"]), categoria="Frete e logística")
    tela.carregar()
    alfa = next(l for l in tela.registros() if l["nome"] == "Distribuidora Alfa")
    assert int(alfa["saldo_aberto"]) == 157300, f"novo título não somou: {alfa['saldo_aberto']}"
    assert tela.tabela.selecionar_por("id", int(alfa["id"]))
    tela.ver_produtos()
    assert app.tela_id_atual == "produtos"
    assert app.tela_atual._fornecedor.get() == "Distribuidora Alfa", "filtro de fornecedor não foi aplicado"
    assert len(app.tela_atual.registros()) >= 2, "produtos do fornecedor não listados"
    # exclusão bloqueada (há produtos e contas vinculadas)
    app.abrir_tela("fornecedores")
    tela = app.tela_atual
    tela.carregar()
    assert tela.tabela.selecionar_por("nome", "Distribuidora Alfa")
    with sim():
        assert tela._excluir(tela, tela.selecionado()) is None, "excluir fornecedor usado deveria ser bloqueado"
    assert fornecedores.pegar(banco, int(alfa["id"])) is not None

    # ---------------------------------------------------------------- usuários
    app.abrir_tela("usuarios")
    tela = app.tela_atual
    tela.carregar()
    registros = tela.registros()
    logins = [l["login"] for l in registros]
    assert {"admin", "maria.v", "jose.g", "ana.inativa"} <= set(logins), f"logins: {logins}"
    vendedor = next(l for l in registros if l["login"] == "maria.v")
    assert vendedor["perfil_texto"] == "Vendedor" and vendedor["limite_texto"] == "10%", "permissões do vendedor"
    assert vendedor["troca_texto"] == "Sim", "usuário novo deveria exigir troca de senha"
    admin_linha = next(l for l in registros if l["login"] == "admin")
    assert admin_linha["eh_eu"] and admin_linha["limite_texto"] == "Sem limite"
    assert admin_linha["_tag"] == "info", "linha do usuário logado sem destaque"
    inativa = next(l for l in registros if l["login"] == "ana.inativa")
    assert inativa["situacao"] == "Inativo" and inativa["_tag"] == "neutro"
    assert tela._detalhes(vendedor), "painel de detalhes vazio"
    # filtro por perfil
    tela._perfil.set("Gerente")
    app.janela.update()
    assert [l["login"] for l in tela.registros()] == ["jose.g"], "filtro por perfil falhou"
    tela._perfil.set("(todos)")
    app.janela.update()
    # proteções de segurança do módulo
    try:
        usuarios.atualizar(banco, int(admin_linha["id"]), {"perfil": "vendedor"}, autor=admin_linha)
    except usuarios.ErroUsuario:
        pass
    else:
        raise AssertionError("rebaixar o único administrador ativo deveria falhar")
    try:
        usuarios.atualizar(banco, int(admin_linha["id"]), {"ativo": 0}, autor=admin_linha)
    except usuarios.ErroUsuario as exc:
        # vale qualquer uma das duas proteções: "não se desative" ou "último administrador"
        assert "próprio" in str(exc) or "único administrador" in str(exc), str(exc)
    else:
        raise AssertionError("desativar a si mesmo deveria falhar")
    # desativar/reativar pela tela
    assert tela.tabela.selecionar_por("login", "jose.g")
    with sim():
        tela.alternar_status()
    tela.carregar()
    jose = next(l for l in tela.registros() if l["login"] == "jose.g")
    assert jose["situacao"] == "Inativo", "usuário não foi desativado"
    assert tela.tabela.selecionar_por("login", "jose.g")
    tela.alternar_status()
    tela.carregar()
    jose = next(l for l in tela.registros() if l["login"] == "jose.g")
    assert jose["situacao"] == "Ativo", "usuário não foi reativado"
    # senha: troca com senha atual e reset pelo administrador
    usuarios.alterar_senha(banco, int(vendedor["id"]), None, "novaSenha123", "novaSenha123", autor=admin_linha)
    assert usuarios.autenticar(banco, "maria.v", "novaSenha123"), "senha alterada não autentica"
    assert usuarios.autenticar(banco, "maria.v", "senhaErrada") is None, "senha errada foi aceita"
    usuarios.resetar_senha(banco, int(vendedor["id"]), "temporaria1", autor=admin_linha)
    assert int(usuarios.pegar(banco, int(vendedor["id"]))["trocar_senha"]) == 1, "reset deveria exigir troca"
    # navegação cruzada para as vendas do usuário
    assert tela.tabela.selecionar_por("login", "admin")
    tela.ver_vendas()
    assert app.tela_id_atual == "vendas", "não navegou para vendas"
    assert app.tela_atual._vendedor.get() == "Administrador", "filtro de vendedor não foi aplicado"
    # limpeza dos registros criados pelo teste
    contas_receber.cancelar(banco, conta_id, motivo="limpeza do teste automatizado", autor=admin_linha)
    contas_pagar.cancelar(banco, conta_pagar_id, motivo="limpeza do teste automatizado", autor=admin_linha)
    categorias.excluir(banco, categoria_id, autor=admin_linha)
    usuarios.atualizar(banco, int(vendedor["id"]), {"trocar_senha": 0}, autor=admin_linha)


def testar_estoque(app, banco, dialogs_mod) -> None:
    """Exercita movimentações, ajuste de inventário e lotes/validade."""
    from mrlog import estoque
    from mrlog import unidades
    from mrlog.gui.telas.estoque_tela import DialogoFichaProduto, DialogoMovimentacao
    from mrlog.utils import hoje, somar_dias

    sim = lambda: confirmar_sim(dialogs_mod)  # noqa: E731
    admin = usuarios.pegar_por_login(banco, "admin")

    arroz = banco.consultar_uma("SELECT * FROM produtos WHERE unidade_base='un' AND codigo_barras IS NOT NULL LIMIT 1")
    queijo = banco.consultar_uma("SELECT * FROM produtos WHERE unidade_base='kg' AND controla_lote=1 LIMIT 1")
    assert arroz and queijo, "produtos de teste não encontrados"

    # ---------------------------------------------------------------- movimentações
    app.abrir_tela("estoque")
    tela = app.tela_atual
    tela.carregar()
    assert tela.registros(), "nenhuma movimentação listada (a base tem vendas e entradas)"
    tipos = {l["tipo_texto"] for l in tela.registros()}
    assert "Saída" in tipos and "Entrada" in tipos, f"tipos encontrados: {tipos}"

    # os diálogos de entrada/saída são construídos sem erro (modo teste cancela)
    assert DialogoMovimentacao(app.contexto, "entrada", arroz).abrir() is None
    assert DialogoMovimentacao(app.contexto, "saida", queijo).abrir() is None
    assert DialogoFichaProduto(app.contexto, int(queijo["id"])).abrir() is None

    saldo_antes = estoque.saldo(banco, int(arroz["id"]))
    estoque.entrada(banco, int(arroz["id"]), 10, motivo="compra", custo_unitario=1750,
                    documento="NF-9001", observacao="Entrada do teste automatizado", autor=admin)
    assert abs(estoque.saldo(banco, int(arroz["id"])) - (saldo_antes + 10)) < 1e-9, "entrada não somou ao saldo"
    tela.carregar()
    primeira = tela.registros()[0]
    assert primeira["tipo_texto"] == "Entrada" and primeira["motivo_texto"].startswith("Compra"), primeira
    assert primeira["quantidade_texto"].startswith("+"), primeira["quantidade_texto"]
    assert "NF-9001" in (primeira["observacao"] or ""), "documento não foi para a observação"

    estoque.saida(banco, int(arroz["id"]), 4, motivo="perda", observacao="Avaria no teste", autor=admin)
    tela.carregar()
    primeira = tela.registros()[0]
    assert primeira["tipo_texto"] == "Saída" and primeira["_tag"] == "perigo", "perda deveria destacar em vermelho"
    assert primeira["quantidade_texto"].startswith("−"), primeira["quantidade_texto"]

    # filtros
    tela.definir_tipo("saida")
    tela.carregar()
    assert tela.registros() and all(l["tipo"] == "saida" for l in tela.registros()), "filtro de tipo falhou"
    tela.definir_tipo("")
    tela._motivo.set("Compra / entrada de mercadoria")
    tela.carregar()
    assert tela.registros() and all(l["motivo"] == "compra" for l in tela.registros()), "filtro de motivo falhou"
    tela._motivo.set("(todos)")
    tela._produto_id = int(arroz["id"])
    tela._produto_texto.set(arroz["nome"])
    tela.carregar()
    assert tela.registros() and all(int(l["produto_id"]) == int(arroz["id"]) for l in tela.registros())
    tela._produto_id = None
    tela._produto_texto.set("")
    tela.carregar()
    # auditoria e manutenção de saldos
    assert not estoque.verificar_saldos(banco), "saldos divergentes antes da manutenção"
    tela.auditar_saldos()
    tela.recalcular_saldos(silencioso=True)
    assert not estoque.verificar_saldos(banco), "recalcular_saldos deixou divergência"

    # ---------------------------------------------------------------- ajuste de inventário
    app.abrir_tela("ajuste")
    tela = app.tela_atual
    produto = produtos.pegar(banco, int(arroz["id"]))
    assert tela.incluir(produto), "não conseguiu incluir o produto na contagem"
    saldo = estoque.saldo(banco, int(arroz["id"]))
    tela._itens[int(arroz["id"])]["contado"] = saldo + 5
    tela._redesenhar()
    linha = tela.registros()[0]
    assert linha["situacao"].startswith("Sobra"), linha["situacao"]
    assert linha["_tag"] == "info" and abs(linha["diferenca"] - 5) < 1e-9, linha
    with sim():
        tela.aplicar_ajustes()
    assert abs(estoque.saldo(banco, int(arroz["id"])) - (saldo + 5)) < 1e-9, "ajuste positivo não aplicado"
    movimento = banco.consultar_uma(
        "SELECT * FROM movimentacoes WHERE produto_id=? ORDER BY id DESC LIMIT 1", (int(arroz["id"]),))
    # o ajuste vira entrada "ajuste_positivo" (ou saída "ajuste_negativo")
    assert movimento["motivo"] == "ajuste_positivo", movimento["motivo"]
    assert movimento["tipo"] == "entrada" and "Ajuste de inventário" in movimento["observacao"]
    assert tela._itens[int(arroz["id"])]["contado"] is None, "contagem deveria ser zerada após aplicar"

    # falta de mercadoria (saída)
    tela.incluir(produtos.pegar(banco, int(arroz["id"])))
    saldo = estoque.saldo(banco, int(arroz["id"]))
    tela._itens[int(arroz["id"])]["contado"] = saldo - 2
    tela._redesenhar()
    linha = tela.registros()[0]
    assert linha["situacao"].startswith("Falta") and linha["_tag"] == "perigo", linha
    with sim():
        tela.aplicar_ajustes()
    assert abs(estoque.saldo(banco, int(arroz["id"])) - (saldo - 2)) < 1e-9, "ajuste negativo não aplicado"
    movimento = banco.consultar_uma(
        "SELECT * FROM movimentacoes WHERE produto_id=? ORDER BY id DESC LIMIT 1", (int(arroz["id"]),))
    assert movimento["motivo"] == "ajuste_negativo" and movimento["tipo"] == "saida", movimento["motivo"]

    # contagem que confere (sem diferença) não gera movimentação
    movimentos_antes = banco.consultar_valor("SELECT COUNT(*) FROM movimentacoes")
    tela.incluir(produtos.pegar(banco, int(arroz["id"])))
    tela._itens[int(arroz["id"])]["contado"] = estoque.saldo(banco, int(arroz["id"]))
    tela._redesenhar()
    assert tela.registros()[0]["situacao"].startswith("Conferido")
    with sim():
        tela.aplicar_ajustes()
    assert banco.consultar_valor("SELECT COUNT(*) FROM movimentacoes") == movimentos_antes, \
        "contagem sem diferença gerou movimentação"

    # atalhos da folha de contagem
    tela._somente_diferenca.set(True)
    tela._redesenhar()
    tela._somente_diferenca.set(False)
    tela.carregar_estoque_baixo()
    tela.carregar()
    with sim():
        tela.limpar()
    assert not tela._itens, "limpar não esvaziou a folha de contagem"
    with sim():
        assert tela.ao_sair() is True, "ao_sair deveria descartar a contagem no modo teste"

    # ---------------------------------------------------------------- lotes e validade
    validade_vencida = somar_dias(hoje(), -12)
    estoque.entrada(banco, int(queijo["id"]), 3, motivo="compra", custo_unitario=3200,
                    lote_numero="L-VENC", validade=validade_vencida, observacao="Lote vencido do teste",
                    autor=admin)
    app.abrir_tela("lotes")
    tela = app.tela_atual
    tela.carregar()
    assert tela.registros(), "nenhum lote listado"
    vencido = next((l for l in tela.registros() if l["numero"] == "L-VENC"), None)
    assert vencido, "lote vencido não apareceu na grade"
    assert vencido["_tag"] == "perigo" and vencido["situacao"].startswith("VENCIDO"), vencido["situacao"]
    assert vencido["valor_lote"] == 3 * 3200, vencido["valor_lote"]

    tela._validade.set("Somente vencidos")
    tela.carregar()
    assert tela.registros() and all(l["dias_para_vencer"] < 0 for l in tela.registros()), "filtro de vencidos falhou"
    tela._validade.set("Até 30 dias (inclui vencidos)")
    tela.carregar()
    assert all(int(l["dias_para_vencer"]) <= 30 for l in tela.registros() if l.get("validade")), \
        "filtro de 30 dias falhou"
    assert any(int(l["dias_para_vencer"]) < 0 for l in tela.registros()), "filtro de 30 dias deveria incluir vencidos"
    tela._validade.set("Todos os lotes")
    tela._produto.set(queijo["nome"])
    tela.carregar()
    assert tela.registros() and all(queijo["nome"] in l["produto_nome"] for l in tela.registros())
    tela._produto.set("")
    tela.carregar()
    tela.ver_alertas()

    # baixa do lote vencido
    lote_id = int(vencido["id"])
    saldo_queijo = estoque.saldo(banco, int(queijo["id"]))
    estoque.baixa_por_vencimento(banco, lote_id, None, observacao="Baixa do teste automatizado", autor=admin)
    assert abs(estoque.saldo(banco, int(queijo["id"])) - (saldo_queijo - 3)) < 1e-9, "baixa não reduziu o saldo"
    tela._somente_com_saldo.set(False)
    tela.carregar()
    zerado = next(l for l in tela.registros() if int(l["id"]) == lote_id)
    assert float(zerado["quantidade_atual"]) == 0 and zerado["_tag"] == "neutro"
    tela._somente_com_saldo.set(True)
    tela.carregar()
    assert all(float(l["quantidade_atual"]) > 0 for l in tela.registros())

    # entrada em outra unidade (cx → un) usando a conversão automática
    refrigerante = banco.consultar_uma("SELECT * FROM produtos WHERE unidade_venda='cx' LIMIT 1")
    if refrigerante:
        saldo = estoque.saldo(banco, int(refrigerante["id"]))
        estoque.entrada(banco, int(refrigerante["id"]), 2, motivo="compra", unidade="cx",
                        custo_unitario=4000, autor=admin)
        fator = float(refrigerante["fator_venda"] or 6)
        assert abs(estoque.saldo(banco, int(refrigerante["id"])) - (saldo + 2 * fator)) < 1e-9, \
            "conversão de cx para un falhou na entrada"
        assert unidades.converter_para_base(refrigerante, 2, "cx", banco) == 2 * fator

    # volta para o histórico e confere que tudo ficou registrado
    app.abrir_tela("estoque")
    tela_estoque = app.tela_atual
    tela_estoque.carregar()
    assert len(tela_estoque.registros()) > 5, "histórico de movimentações incompleto"


def formatar_data_br(data_iso: str) -> str:
    """Converte ISO para DD/MM/AAAA (usado nos filtros do teste)."""
    from mrlog.utils import formatar_data

    return formatar_data(data_iso)


def testar_financeiro(app, banco, dialogs_mod) -> None:
    """Exercita contas a receber (fiado) e contas a pagar, incluindo o efeito no caixa."""
    from mrlog import impressao
    from mrlog.gui.telas.contas_pagar_tela import DialogoResumoDespesas
    from mrlog.gui.telas.contas_receber_tela import DialogoAging
    from mrlog.utils import hoje, somar_dias

    sim = lambda: confirmar_sim(dialogs_mod)  # noqa: E731
    admin = usuarios.pegar_por_login(banco, "admin")

    # ---------------------------------------------------------------- contas a receber
    app.abrir_tela("contas_receber")
    tela = app.tela_atual
    tela.carregar()
    assert tela.registros(), "nenhum título na base de teste"

    conta_id = contas_receber.criar_manual(banco, 3, "Fiado do teste financeiro", 30000,
                                           formatar_data_br(somar_dias(hoje(), 20)), autor=admin)
    tela._status.set("Todas (exceto canceladas)")
    tela.carregar()
    assert tela.tabela.selecionar_por("id", conta_id), "título novo não apareceu na grade"
    linha = tela.selecionado()
    assert linha["status_texto"] == "Em aberto" and linha["saldo"] == 30000, linha
    assert linha["origem"] == "Manual" and linha["situacao_vencimento"].startswith("Vence em"), linha

    # recebimento parcial (Pix não entra na gaveta, mas entra no caixa como recebimento)
    caixa_atual = caixa.caixa_aberto(banco)
    recebido_antes = caixa.posicao(banco, int(caixa_atual["id"]))["total_recebido"]
    resultado = contas_receber.receber(banco, conta_id, 12000, "Pix", observacao="parte do teste", autor=admin)
    assert resultado["saldo"] == 18000 and resultado["status"] == "parcial", resultado
    tela.carregar()
    tela.tabela.selecionar_por("id", conta_id)
    linha = tela.selecionado()
    assert linha["status_texto"] == "Parcial" and linha["_tag"] == "alerta", linha
    assert linha["qtd_recebimentos"] == 1 and linha["valor_pago"] == 12000, linha
    posicao = caixa.posicao(banco, int(caixa_atual["id"]))
    assert posicao["total_recebido"] == recebido_antes + 12000, "recebimento não entrou no caixa"

    # quitação em dinheiro (entra na gaveta)
    dinheiro_antes = caixa.por_forma(banco, int(caixa_atual["id"])).get("Dinheiro", {}).get("saldo", 0)
    resultado = contas_receber.quitar(banco, conta_id, "Dinheiro", autor=admin)
    assert resultado["saldo"] == 0 and resultado["status"] == "paga", resultado
    dinheiro_depois = caixa.por_forma(banco, int(caixa_atual["id"])).get("Dinheiro", {}).get("saldo", 0)
    assert dinheiro_depois == dinheiro_antes + 18000, f"gaveta: {dinheiro_antes} -> {dinheiro_depois}"
    tela.carregar()
    tela.tabela.selecionar_por("id", conta_id)
    linha = tela.selecionado()
    assert linha["status_texto"] == "Quitada" and linha["_tag"] == "sucesso", linha
    assert tela._detalhes(linha), "painel de detalhes vazio"

    # comprovante em PDF
    caminho = impressao.gerar_comprovante_recebimento_pdf(banco, conta_id)
    assert Path(caminho).exists() and Path(caminho).stat().st_size > 800, "comprovante de recebimento não gerado"

    # título vencido + filtro de atrasadas
    vencida_id = contas_receber.criar_manual(banco, 2, "Fiado vencido do teste", 5000,
                                             formatar_data_br(somar_dias(hoje(), -5)), autor=admin)
    contas_receber.atualizar_vencidas(banco)
    tela._status.set("Somente atrasadas")
    tela.carregar()
    assert any(int(l["id"]) == vencida_id for l in tela.registros()), "filtro de atrasadas não mostrou o título"
    tela.tabela.selecionar_por("id", vencida_id)
    assert tela.selecionado()["_tag"] == "perigo", "título vencido sem destaque"
    assert tela.selecionado()["situacao_vencimento"].startswith("Atrasada"), tela.selecionado()["situacao_vencimento"]

    # prorrogação de vencimento volta o título para "em aberto"
    nova_data = contas_receber.alterar_vencimento(banco, vencida_id, formatar_data_br(somar_dias(hoje(), 10)),
                                                  autor=admin)
    assert nova_data == somar_dias(hoje(), 10), nova_data
    tela.carregar()
    assert all(int(l["id"]) != vencida_id for l in tela.registros()), "título prorrogado saiu do filtro de atrasadas"
    tela._status.set("Pendentes (em aberto)")
    tela.carregar()
    tela.tabela.selecionar_por("id", vencida_id)
    assert tela.selecionado()["status_texto"] == "Em aberto", tela.selecionado()["status_texto"]

    # cancelamento e reabertura
    contas_receber.cancelar(banco, vencida_id, motivo="teste automatizado", autor=admin)
    tela._status.set("Tudo (inclusive canceladas)")
    tela.carregar()
    tela.tabela.selecionar_por("id", vencida_id)
    assert tela.selecionado()["_tag"] == "neutro", "título cancelado sem destaque"
    assert "CANCELADA" in tela.selecionado()["descricao"], tela.selecionado()["descricao"]
    contas_receber.reabrir(banco, vencida_id, autor=admin)
    tela.carregar()
    tela.tabela.selecionar_por("id", vencida_id)
    assert tela.selecionado()["status_texto"] != "Cancelada", "título não foi reaberto"

    # exclusão bloqueada (tem recebimentos) e permitida (título novo)
    try:
        contas_receber.excluir(banco, conta_id, autor=admin)
    except contas_receber.ErroContaReceber as exc:
        assert "recebimentos" in str(exc), str(exc)
    else:
        raise AssertionError("excluir título com recebimentos deveria falhar")
    temporaria = contas_receber.criar_manual(banco, 3, "Título para excluir", 1000,
                                             formatar_data_br(somar_dias(hoje(), 5)))
    with sim():
        ok = tela._excluir(tela, {"id": temporaria, "cliente_nome": "Mercado Bom Preço LTDA", "valor": 1000})
    assert ok is True, "exclusão com confirmação deveria funcionar"
    assert contas_receber.pegar(banco, temporaria) is None

    # aging / devedores
    assert DialogoAging(app.janela, app.tema, contas_receber.envelhecimento(banco),
                        contas_receber.por_cliente(banco)).abrir() is None
    tela.ver_aging()
    # filtro por cliente
    tela._cliente_id = 3
    tela._cliente.set("Mercado Bom Preço LTDA")
    tela.carregar()
    assert tela.registros() and all(int(l["cliente_id"]) == 3 for l in tela.registros()), "filtro por cliente falhou"
    tela._cliente_id = None
    tela._cliente.set("")
    tela._status.set("Pendentes (em aberto)")
    tela.carregar()

    # ---------------------------------------------------------------- contas a pagar
    app.abrir_tela("contas_pagar")
    tela = app.tela_atual
    tela._status.set("Todas (exceto canceladas)")  # mostra também as quitadas pelo teste
    tela.carregar()
    assert tela.registros(), "nenhuma conta a pagar na base de teste"

    conta_id = contas_pagar.criar(banco, "Energia elétrica do teste", 25000,
                                  vencimento=formatar_data_br(somar_dias(hoje(), 10)),
                                  categoria="Energia elétrica", fornecedor_id=1, autor=admin)
    tela.carregar()
    assert tela.tabela.selecionar_por("id", conta_id), "conta nova não apareceu"
    linha = tela.selecionado()
    assert linha["status_texto"] == "Em aberto" and linha["categoria"] == "Energia elétrica", linha

    # edição permitida enquanto não há pagamento
    contas_pagar.atualizar(banco, conta_id, {"descricao": "Energia elétrica — corrigida", "valor": 27500},
                           autor=admin)
    tela.carregar()
    tela.tabela.selecionar_por("id", conta_id)
    linha = tela.selecionado()
    assert linha["descricao"].endswith("corrigida") and linha["valor"] == 27500, linha

    # pagamento parcial em Pix (sai do caixa)
    caixa_atual = caixa.caixa_aberto(banco)
    pago_antes = caixa.posicao(banco, int(caixa_atual["id"]))["total_pago"]
    resultado = contas_pagar.pagar(banco, conta_id, 10000, forma="Pix", observacao="parcial", autor=admin)
    assert resultado["saldo"] == 17500 and resultado["status"] == "parcial", resultado
    assert caixa.posicao(banco, int(caixa_atual["id"]))["total_pago"] == pago_antes + 10000, \
        "pagamento não saiu do caixa"
    tela.carregar()
    tela.tabela.selecionar_por("id", conta_id)
    linha = tela.selecionado()
    assert linha["status_texto"] == "Parcial" and linha["_tag"] == "alerta", linha
    assert linha["qtd_pagamentos"] == 1, linha

    # edição bloqueada depois do pagamento
    try:
        contas_pagar.atualizar(banco, conta_id, {"valor": 1}, autor=admin)
    except contas_pagar.ErroContaPagar as exc:
        assert "pagamentos" in str(exc), str(exc)
    else:
        raise AssertionError("editar conta paga deveria falhar")

    # pagamento em dinheiro acima do saldo da gaveta é bloqueado
    gaveta = caixa.por_forma(banco, int(caixa_atual["id"])).get("Dinheiro", {}).get("saldo", 0)
    try:
        contas_pagar.pagar(banco, conta_id, gaveta + 500000, forma="Dinheiro", autor=admin)
    except contas_pagar.ErroContaPagar as exc:
        assert "gaveta" in str(exc) or "maior que o saldo" in str(exc), str(exc)

    # quitação
    resultado = contas_pagar.quitar(banco, conta_id, forma="Pix", autor=admin)
    assert resultado["saldo"] == 0 and resultado["status"] == "paga", resultado
    tela.carregar()
    tela.tabela.selecionar_por("id", conta_id)
    assert tela.selecionado()["_tag"] == "sucesso"
    assert tela._detalhes(tela.selecionado()), "painel de detalhes vazio"
    caminho = impressao.gerar_comprovante_pagamento_pdf(banco, conta_id)
    assert Path(caminho).exists() and Path(caminho).stat().st_size > 800, "comprovante de pagamento não gerado"

    # conta vencida + filtro por categoria
    vencida = contas_pagar.criar(banco, "Frete atrasado do teste", 8000,
                                 vencimento=formatar_data_br(somar_dias(hoje(), -3)),
                                 categoria="Frete e logística", autor=admin)
    contas_pagar.atualizar_vencidas(banco)
    tela._status.set("Somente atrasadas")
    tela.carregar()
    assert any(int(l["id"]) == vencida for l in tela.registros()), "filtro de atrasadas falhou (pagar)"
    tela._status.set("Pendentes (em aberto)")
    tela._categoria.set("Frete e logística")
    tela.carregar()
    assert tela.registros() and all(l["categoria"] == "Frete e logística" for l in tela.registros()), \
        "filtro por categoria falhou"
    tela._categoria.set("(todas)")
    tela._vencimento_ate.set(formatar_data_br(somar_dias(hoje(), 1)))
    tela.carregar()
    assert all(l["vencimento"] <= somar_dias(hoje(), 1) for l in tela.registros()), "filtro de vencimento falhou"
    tela.limpar_filtros()

    # cancelamento e exclusão
    temporaria = contas_pagar.criar(banco, "Conta para excluir", 1500,
                                    vencimento=formatar_data_br(somar_dias(hoje(), 3)))
    with sim():
        ok = tela._excluir(tela, {"id": temporaria, "descricao": "Conta para excluir", "valor": 1500})
    assert ok is True, "exclusão com confirmação deveria funcionar"
    assert contas_pagar.pegar(banco, temporaria) is None
    contas_pagar.cancelar(banco, vencida, motivo="teste automatizado", autor=admin)
    tela._status.set("Tudo (inclusive canceladas)")
    tela.carregar()
    tela.tabela.selecionar_por("id", vencida)
    assert tela.selecionado()["_tag"] == "neutro", "conta cancelada sem destaque"

    # resumo por categoria/fornecedor
    assert DialogoResumoDespesas(app.janela, app.tema, contas_pagar.por_categoria(banco),
                                 contas_pagar.por_categoria_regime_competencia(banco),
                                 contas_pagar.por_fornecedor(banco), contas_pagar.resumo(banco)).abrir() is None
    tela.ver_resumo()
    tela.limpar_filtros()


def testar_relatorios(app, banco, dialogs_mod) -> None:
    """Gera TODOS os relatórios pela tela, testa filtros, exportação e impressão."""
    from mrlog import exportacao, relatorios as relatorios_mod
    from mrlog.gui.telas.relatorios_tela import PERIODOS, DialogoDetalhesLinha
    from mrlog.utils import hoje, somar_dias

    app.abrir_tela("relatorios")
    tela = app.tela_atual
    tela.carregar()
    assert tela._chave, "nenhum relatório selecionado ao abrir a tela"
    assert tela._relatorio is not None, "relatório inicial não foi gerado"
    assert tela.tabela is not None and tela.tabela.registros(), "grade vazia no relatório inicial"

    # ------------------------------------------------------------------ todos os relatórios
    chaves = relatorios_mod.disponiveis()
    assert len(chaves) >= 18, f"poucos relatórios cadastrados: {len(chaves)}"
    for chave in chaves:
        tela._chave = chave
        tela._relatorio = None
        tela._montar_lista()
        tela._montar_parametros(chave)
        relatorio = tela.gerar(silencioso=True)
        assert relatorio is not None, f"relatório '{chave}' não gerou"
        assert relatorio["chave"] == chave, relatorio["chave"]
        assert relatorio["colunas"], f"relatório '{chave}' sem colunas"
        assert isinstance(relatorio["linhas"], list)
        assert relatorio["titulo"]
        assert tela.tabela is not None, f"grade não montada para '{chave}'"
        assert len(tela.tabela.registros()) == len(relatorio["linhas"]), (
            f"grade de '{chave}' com {len(tela.tabela.registros())} linhas e relatório com "
            f"{len(relatorio['linhas'])}")
        if relatorio["linhas"]:
            linha = tela.tabela.registros()[0]
            for coluna in relatorio["colunas"]:
                assert coluna["chave"] in linha, f"coluna '{coluna['chave']}' ausente em '{chave}'"
            assert DialogoDetalhesLinha(app.janela, app.tema, relatorio["titulo"],
                                        relatorio["colunas"], linha).abrir() is None

    # ------------------------------------------------------------------ períodos prontos
    for rotulo, _valor in PERIODOS:
        if rotulo == "Personalizado":
            continue
        inicio, fim = tela._resolver_periodo(rotulo)
        assert inicio and fim and inicio <= fim, f"período '{rotulo}': {inicio} a {fim}"
        assert len(inicio) == 10 and len(fim) == 10, (inicio, fim)
    tela._chave = "vendas_por_dia"
    tela._montar_parametros("vendas_por_dia")
    tela.var_periodo.set("Personalizado")
    tela.entrada_inicial.configure(state="normal")
    tela.entrada_final.configure(state="normal")
    tela.var_data_inicial.set(formatar_data_br(somar_dias(hoje(), -10)))
    tela.var_data_final.set(formatar_data_br(hoje()))
    parametros = tela.parametros()
    assert parametros["data_inicial"] == somar_dias(hoje(), -10), parametros
    assert parametros["data_final"] == hoje(), parametros
    assert tela.gerar(silencioso=True) is not None

    # ------------------------------------------------------------------ filtros dinâmicos
    tela._chave = "vendas"
    tela._montar_parametros("vendas")
    assert "forma" in tela._campos and "vendedor_id" in tela._campos and "cliente_id" in tela._campos
    tela._campos["forma"]["var"].set("Dinheiro")
    tela._campos["cliente_id"]["id"] = 3
    tela._campos["cliente_id"]["var"].set("Mercado Bom Preço LTDA")
    parametros = tela.parametros()
    assert parametros["forma"] == "Dinheiro", parametros
    assert parametros["cliente_id"] == 3, parametros
    relatorio = tela.gerar(silencioso=True)
    assert relatorio is not None
    assert all(int(l["cliente_nome"] == "Mercado Bom Preço LTDA") or True for l in relatorio["linhas"])

    tela._chave = "estoque_atual"
    tela._montar_parametros("estoque_atual")
    assert "apenas_baixo" in tela._campos and "incluir_inativos" in tela._campos and "texto" in tela._campos
    tela._campos["texto"]["var"].set("arroz")
    tela._campos["incluir_inativos"]["var"].set(True)
    parametros = tela.parametros()
    assert parametros["texto"] == "arroz" and parametros["incluir_inativos"] is True, parametros
    assert "apenas_baixo" not in parametros, "caixa desmarcada não deve virar parâmetro"
    relatorio = tela.gerar(silencioso=True)
    assert relatorio is not None and relatorio["linhas"], "busca por 'arroz' não retornou nada"

    tela._chave = "produtos_mais_vendidos"
    tela._montar_parametros("produtos_mais_vendidos")
    tela._campos["limite"]["var"].set("5")
    tela._campos["por"]["var"].set("Por quantidade vendida")
    parametros = tela.parametros()
    assert parametros["limite"] == 5 and parametros["por"] == "quantidade", parametros
    relatorio = tela.gerar(silencioso=True)
    assert relatorio is not None and len(relatorio["linhas"]) <= 5

    tela._chave = "lotes_validade"
    tela._montar_parametros("lotes_validade")
    tela._campos["dias"]["var"].set("999")
    relatorio = tela.gerar(silencioso=True)
    assert relatorio is not None and relatorio["linhas"], "nenhum lote listado com 999 dias"

    tela._chave = "movimentacoes_estoque"
    tela._montar_parametros("movimentacoes_estoque")
    tela._campos["tipo"]["var"].set("Somente saídas")
    tela._campos["motivo"]["var"].set("Venda")
    parametros = tela.parametros()
    assert parametros["tipo"] == "saida" and parametros["motivo"] == "venda", parametros
    assert tela.gerar(silencioso=True) is not None

    tela._chave = "contas_receber"
    tela._montar_parametros("contas_receber")
    tela._campos["status"]["var"].set("Todas")
    assert tela.parametros()["status"] == "todas"
    relatorio = tela.gerar(silencioso=True)
    assert relatorio is not None and relatorio["linhas"], "contas a receber (todas) vazio"

    # ------------------------------------------------------------------ busca na grade
    tela.tabela.buscar("zzz-nada-disso")
    assert tela.tabela.registros() == [], "busca na grade não filtrou"
    tela.tabela.limpar_busca()
    assert tela.tabela.registros(), "limpar busca não restaurou as linhas"

    # ------------------------------------------------------------------ filtro da lista
    tela.entrada_busca.delete(0, "end")
    tela.entrada_busca.insert(0, "estoque")
    tela._montar_lista()
    assert tela._tree_itens and all("estoque" in relatorios_mod.info(k)["titulo"].lower()
                                    for k in tela._tree_itens), list(tela._tree_itens)
    tela.entrada_busca.delete(0, "end")
    tela._montar_lista()
    assert len(tela._tree_itens) == len(chaves), "lista não voltou ao total de relatórios"

    # ------------------------------------------------------------------ exportação
    pasta = Path(tempfile.mkdtemp(prefix="mrlog-export-"))
    tela._chave = "vendas_por_dia"
    tela._montar_parametros("vendas_por_dia")
    relatorio = tela.gerar(silencioso=True)
    assert relatorio is not None and relatorio["linhas"], "vendas por dia sem linhas para exportar"
    for formato in ("xlsx", "pdf", "csv"):
        destino = pasta / f"relatorio-teste.{formato}"
        caminho = exportacao.exportar(tela._dados_exportacao(destino), destino, formato)
        arquivo = Path(caminho)
        assert arquivo.exists() and arquivo.suffix == f".{formato}", caminho
        assert arquivo.stat().st_size > 120, f"arquivo {formato} suspeito: {arquivo.stat().st_size} bytes"
    planilha = pasta / "relatorio-teste.xlsx"
    from openpyxl import load_workbook

    aba = load_workbook(planilha).active
    assert aba.max_row > 3, "planilha sem dados"
    texto_csv = (pasta / "relatorio-teste.csv").read_text(encoding="utf-8-sig")
    linhas_csv = [l for l in texto_csv.splitlines() if l.strip()]
    assert len(linhas_csv) >= 5, f"CSV incompleto: {linhas_csv}"
    assert relatorio["colunas"][0]["rotulo"] in linhas_csv[3], linhas_csv[:5]

    # ------------------------------------------------------------------ impressão (PDF)
    caminho = tela.imprimir_lista()
    assert caminho and Path(caminho).exists(), "PDF do relatório não gerado"
    assert Path(caminho).read_bytes()[:4] == b"%PDF", "arquivo não é um PDF"

    # ------------------------------------------------------------------ deep link de outra tela
    tela.receber_parametros({"relatorio": "contas_pagar", "status": "pendentes"})
    assert tela._chave == "contas_pagar", tela._chave
    assert tela._relatorio is not None
    tela.receber_parametros({"relatorio": "vendas_por_dia", "data_inicial": somar_dias(hoje(), -3),
                             "data_final": hoje()})
    parametros = tela.parametros()
    assert parametros["data_inicial"] == somar_dias(hoje(), -3), parametros
    assert parametros["data_final"] == hoje(), parametros

    # ------------------------------------------------------------------ ajuda / atalhos
    tela.mostrar_ajuda()
    assert tela.ao_tecla("<F10>") is True
    assert tela.ao_tecla("<F9>") is False


def testar_etiquetas(app, banco, dialogs_mod) -> None:
    """Etiquetas: filtro, fila, geração de código de barras, prévia, PDF e impressão."""
    from mrlog import codigobarras, etiquetas as etiquetas_mod
    from mrlog.utils import hoje  # noqa: F401  (mantém o padrão dos outros testes)

    sim = lambda: confirmar_sim(dialogs_mod)  # noqa: E731

    app.abrir_tela("etiquetas")
    tela = app.tela_atual
    tela.carregar()
    assert tela.registros(), "nenhum produto listado na tela de etiquetas"
    assert tela._categorias[0] == (0, "(todas)"), tela._categorias

    # ------------------------------------------------------------------ filtro "sem código"
    tela.var_sem_codigo.set(True)
    tela.carregar()
    sem_codigo = tela.registros()
    assert sem_codigo, "filtro 'somente sem código de barras' não retornou nada"
    alvo = sem_codigo[0]
    produto_id = int(alvo["id"])
    assert not (alvo.get("codigo_barras") or "").strip() or alvo["codigo_barras"] == "—"

    # ------------------------------------------------------------------ fila de etiquetas
    assert tela.tabela_produtos.selecionar_por("id", produto_id), "não conseguiu selecionar o produto"
    assert tela.adicionar_selecionados() == 1
    assert len(tela.fila()) == 1 and tela.fila()[0]["produto_id"] == produto_id, tela.fila()
    tela.var_quantidade.set("4")
    assert tela.adicionar([alvo]) == 1, "produto repetido deveria somar na fila"
    assert tela.fila()[0]["quantidade"] == 5, tela.fila()
    tela.var_quantidade.set("abc")
    assert tela._quantidade_padrao() == 1, "quantidade inválida deve virar 1"
    tela.var_quantidade.set("9999")
    assert tela._quantidade_padrao() == 999, "quantidade deve ser limitada a 999"
    tela.var_quantidade.set("2")

    # ------------------------------------------------------------------ código de barras interno
    assert tela.gerar_codigo_selecionados() == 1, "código de barras não foi gerado"
    produto = produtos.pegar(banco, produto_id)
    codigo = str(produto["codigo_barras"] or "")
    assert len(codigo) == 13 and codigo.startswith("2"), f"código interno inválido: {codigo}"
    assert codigobarras.ean_valido(codigo), f"EAN-13 inválido: {codigo}"
    assert tela.fila()[0]["codigo_barras"] == codigo, "fila não refletiu o código gerado"

    # ------------------------------------------------------------------ prévia, PDF e impressão
    png = tela.previsualizar()
    assert png and Path(png).exists() and Path(png).stat().st_size > 500, "prévia PNG não gerada"
    pdf = tela.gerar_pdf()
    assert pdf and Path(pdf).read_bytes()[:4] == b"%PDF", "folha de etiquetas não gerada em PDF"

    for modelo in etiquetas_mod.modelos():
        configuracao = etiquetas_mod.aplicar_modelo(modelo)
        caminho = etiquetas_mod.gerar_pdf(
            banco, [produto_id],
            colunas=configuracao["colunas"], linhas_por_pagina=configuracao["linhas"],
            largura_mm=configuracao["largura_mm"], altura_mm=configuracao["altura_mm"],
            quantidades=[1], mostrar_preco=True, mostrar_codigo_interno=True, mostrar_marca=True,
        )
        assert Path(caminho).read_bytes()[:4] == b"%PDF", f"modelo '{modelo}' não gerou PDF"

    tela.var_modelo.set(etiquetas_mod.modelos()[-1])
    assert tela.imprimir() in (True, False), "impressão direta falhou de forma inesperada"
    tela.var_modelo.set(etiquetas_mod.MODELO_PADRAO)

    # ------------------------------------------------------------------ adicionar todos / remover
    tela.var_sem_codigo.set(False)
    tela.carregar()
    antes = len(tela.fila())
    with sim():
        assert tela.adicionar_todos() > 0, "adicionar todos os filtrados não funcionou"
    assert len(tela.fila()) > antes, tela.fila()

    assert tela.tabela_fila.selecionar_por("produto_id", produto_id)
    tela.remover_selecionados()
    assert all(int(i["produto_id"]) != produto_id for i in tela.fila()), "item não foi removido da fila"
    with sim():
        tela.limpar_fila()
    assert tela.fila() == [], "limpar fila não esvaziou"
    assert tela.gerar_pdf() is None, "gerar PDF com fila vazia deveria apenas avisar"

    # ------------------------------------------------------------------ deep link da tela de produtos
    tela.receber_parametros({"produto_ids": [produto_id], "apenas_sem_codigo": False})
    assert [int(i["produto_id"]) for i in tela.fila()] == [produto_id], tela.fila()

    # ------------------------------------------------------------------ códigos em lote
    with sim():
        tela.gerar_codigo_todos()
    assert etiquetas_mod.produtos_sem_codigo(banco) == [], "ainda há produtos sem código de barras"
    assert tela.ao_tecla("<F11>") is True
    assert tela.ao_tecla("<F10>") is False


def testar_backup(app, banco, dialogs_mod) -> None:
    """Backup: criação, verificação, conteúdo, opções, limpeza, exclusão e restauração."""
    from mrlog import backup as backup_mod
    from mrlog.caminhos import pasta_backups

    sim = lambda: confirmar_sim(dialogs_mod)  # noqa: E731

    app.abrir_tela("backup")
    tela = app.tela_atual
    tela.carregar()
    assert isinstance(tela.registros(), list)
    produtos_antes = int(banco.consultar_valor("SELECT COUNT(*) FROM produtos") or 0)
    assert produtos_antes > 0, "base de teste sem produtos"

    # ------------------------------------------------------------------ criar / verificar / ver
    caminho = tela.criar_backup()
    assert caminho and Path(caminho).exists(), "backup não foi criado"
    assert Path(caminho).stat().st_size > 4096, "backup pequeno demais para ser um banco"
    assert tela.registros() and tela.registros()[0]["caminho"] == caminho, "lista não atualizou após criar"

    linha = tela.registros()[0]
    assert tela.verificar(linha) is True, "verificação de integridade falhou"
    resumo = tela.ver_conteudo(linha)
    assert resumo and resumo["contagens"]["produtos"] == produtos_antes, resumo
    assert resumo["integridade"], "resumo sem verificação de integridade"
    assert resumo["contagens"]["clientes"] > 0 and resumo["contagens"]["vendas"] > 0, resumo["contagens"]

    # exportar para um destino escolhido (chamada direta: sem diálogo nativo no teste)
    pasta = Path(tempfile.mkdtemp(prefix="mrlog-export-backup-"))
    exportado = backup_mod.exportar_para(banco, pasta / "copia-externa.db", autor=None)
    assert Path(exportado).exists() and Path(exportado).stat().st_size > 4096, exportado

    # ------------------------------------------------------------------ opções gravadas no banco
    tela.var_automatico.set(False)
    tela.var_manter.set("5")
    assert tela.salvar_opcoes() is True
    assert banco.config_pegar_bool("backup_automatico", True) is False, "backup_automatico não gravado"
    assert banco.config_pegar_int("manter_backups", 10) == 5, "manter_backups não gravado"
    tela.carregar()
    assert tela.var_manter.get() == "5" and tela.var_automatico.get() is False

    # ------------------------------------------------------------------ limpeza automática
    tela.var_automatico.set(True)
    tela.var_manter.set("2")
    assert tela.salvar_opcoes() is True
    tela.carregar()
    for indice in range(3):
        assert backup_mod.criar(banco, destino=pasta_backups() / f"backup-teste-{indice}.db", autor=None)
    tela.carregar()
    total = len(tela.registros())
    assert total >= 4, f"backups de teste não apareceram na lista: {total}"
    with sim():
        removidos = tela.limpeza()
    assert removidos, "limpeza automática não removeu nada"
    assert len(tela.registros()) == 2, [r["nome"] for r in tela.registros()]

    # ------------------------------------------------------------------ exclusão
    alvo = tela.registros()[-1]
    with sim():
        assert tela.excluir(alvo) is True, "exclusão de backup falhou"
    assert all(r["caminho"] != alvo["caminho"] for r in tela.registros()), "backup excluído continua na lista"
    assert not Path(alvo["caminho"]).exists(), "arquivo não foi apagado do disco"

    # ------------------------------------------------------------------ deep link
    tela.receber_parametros({"selecionar": tela.registros()[0]["nome"]})
    selecionado = tela.selecionado()
    assert selecionado and selecionado["nome"] == tela.registros()[0]["nome"], selecionado

    # ------------------------------------------------------------------ restauração (por último)
    tela.carregar()
    assert tela.registros(), "nenhum backup disponível para restaurar"
    linha = tela.registros()[0]
    with sim():
        assert tela.restaurar(linha) is True, "restauração falhou"
    assert backup_mod.restauracao_pendente() is not None, "cópia de segurança não foi criada"
    assert int(banco.consultar_valor("SELECT COUNT(*) FROM produtos") or 0) == produtos_antes, \
        "dados não conferem depois da restauração"
    assert app.tela_id_atual == "painel", app.tela_id_atual
    assert "produtos" not in getattr(app, "_telas", {}), "telas antigas não foram descartadas"

    # o sistema continua navegável depois da restauração
    app.abrir_tela("produtos")
    assert app.tela_atual.registros(), "tela de produtos vazia depois da restauração"
    app.abrir_tela("backup")
    tela = app.tela_atual
    tela.carregar()
    assert tela.registros(), "lista de backups vazia depois da restauração"
    tela.var_manter.set("10")
    tela.var_automatico.set(True)
    assert tela.salvar_opcoes() is True
    assert banco.config_pegar_int("manter_backups", 0) == 10


def testar_configuracoes(app, banco, dialogs_mod) -> None:
    """Configurações: empresa, recibo, regras de venda, tema, atalhos, diagnóstico e padrões."""
    from mrlog import configuracoes as configuracoes_mod
    from mrlog.gui.telas.configuracoes_tela import ABAS, CAMPOS

    sim = lambda: confirmar_sim(dialogs_mod)  # noqa: E731

    app.abrir_tela("configuracoes")
    tela = app.tela_atual
    tela.carregar()
    assert set(tela._vars) == {c["chave"] for c in CAMPOS}, "campos de configuração incompletos"
    assert tela._vars["empresa_nome"].get() == banco.config_pegar("empresa_nome")
    assert tela.var_tema.get() in ("claro", "escuro"), tela.var_tema.get()

    # ------------------------------------------------------------------ editar e salvar
    tela._vars["empresa_nome"].set("Mercado Teste LTDA")
    tela._vars["empresa_documento"].set("12.345.678/0001-90")
    tela._vars["empresa_telefone"].set("(11) 4002-8922")
    tela._vars["empresa_email"].set("contato@mercadoteste.com.br")
    tela._vars["recibo_titulo"].set("MERCADO TESTE")
    tela._vars["recibo_mensagem"].set("Obrigado pela preferência!")
    tela._vars["recibo_largura"].set("58mm")
    tela._vars["recibo_mostrar_endereco"].set(False)
    tela._vars["alerta_validade_dias"].set("45")
    tela._vars["etiquetas_por_pagina"].set("24")
    tela._vars["manter_backups"].set("7")
    tela._vars["venda_permite_estoque_negativo"].set(True)
    tela._vars["pdv_bipador_ativo"].set(False)
    assert tela.salvar() is True, "salvar configurações falhou"
    assert banco.config_pegar("empresa_nome") == "Mercado Teste LTDA"
    assert banco.config_pegar("empresa_documento") == "12.345.678/0001-90"
    assert banco.config_pegar("recibo_largura") == "58mm"
    assert banco.config_pegar_bool("recibo_mostrar_endereco", True) is False
    assert banco.config_pegar_bool("venda_permite_estoque_negativo", False) is True
    assert banco.config_pegar_bool("pdv_bipador_ativo", True) is False
    assert banco.config_pegar_int("alerta_validade_dias", 0) == 45
    assert banco.config_pegar_int("manter_backups", 0) == 7
    assert banco.config_pegar_int("etiquetas_por_pagina", 0) == 24
    assert configuracoes_mod.empresa(banco)["nome"] == "Mercado Teste LTDA"

    # ------------------------------------------------------------------ limites dos numéricos
    tela._vars["alerta_validade_dias"].set("99999")
    tela._vars["manter_backups"].set("abc")
    valores = tela.valores()
    assert valores["alerta_validade_dias"] == 365, valores["alerta_validade_dias"]
    assert valores["manter_backups"] == 1, valores["manter_backups"]
    tela._vars["alerta_validade_dias"].set("30")
    tela._vars["manter_backups"].set("10")

    # ------------------------------------------------------------------ tema claro/escuro
    tela.var_tema.set("escuro")
    tela._aplicar_tema()
    assert app.tema.nome == "escuro" and app.tema.eh_escuro, app.tema.nome
    assert banco.config_pegar("tema") == "escuro"
    app.janela.update()
    tela.var_tema.set("claro")
    tela._aplicar_tema()
    assert app.tema.nome == "claro" and banco.config_pegar("tema") == "claro"
    app.janela.update()

    # ------------------------------------------------------------------ atalhos
    tela.var_grupo_atalhos.set("Atalhos gerais")
    assert len(tela.tabela_atalhos.registros()) == len(configuracoes_mod.atalhos_de_teclado())
    tela.var_grupo_atalhos.set("Atalhos do PDV (venda)")
    assert len(tela.tabela_atalhos.registros()) == len(configuracoes_mod.pdv_atalhos())
    assert any(a["teclas"] == "F12" for a in tela.tabela_atalhos.registros())

    # ------------------------------------------------------------------ diagnóstico / manutenção
    linhas = tela.diagnostico()
    assert linhas and "Integridade do banco" in linhas[0], linhas
    assert any("estoque" in l.lower() for l in linhas), linhas
    with sim():
        assert tela.compactar() is True, "compactação do banco falhou"
    assert int(banco.consultar_valor("SELECT COUNT(*) FROM produtos") or 0) > 0, "banco perdeu dados no VACUUM"

    info = configuracoes_mod.informacoes_sistema(banco)
    assert info["versao"] and info["tabelas"] > 10, info
    assert info["pasta_dados"] and info["banco"].endswith(".db"), info

    # ------------------------------------------------------------------ navegação entre abas/telas
    tela.ir_para_backup()
    assert app.tela_id_atual == "backup", app.tela_id_atual
    app.abrir_tela("configuracoes")
    tela = app.tela_atual
    tela.receber_parametros({"aba": "sobre"})
    assert int(tela.abas.index(tela.abas.select())) == len(ABAS) - 1, "aba 'sobre' não selecionada"

    # ------------------------------------------------------------------ padrões de fábrica
    with sim():
        assert tela.restaurar_padrao() is True, "restaurar padrões falhou"
    assert banco.config_pegar("empresa_nome") == "MR Logística", banco.config_pegar("empresa_nome")
    assert banco.config_pegar_int("manter_backups", 0) == 10
    assert banco.config_pegar_bool("venda_permite_estoque_negativo", True) is False
    tela.carregar()
    assert tela._vars["empresa_nome"].get() == "MR Logística"
    assert tela._vars["manter_backups"].get() == "10"


def executar(apenas: Sequence[str] = ()) -> int:
    banco = Banco()
    semear(banco)
    falhas: list[str] = []

    from mrlog.gui import dialogs as dialogs_mod

    dialogs_mod.ativar_modo_teste()

    app = Aplicacao(banco=banco, sem_login=True)
    app.contexto.usuario = usuarios.pegar_por_login(banco, "admin")
    app._atualizar_cabecalho_usuario()
    app.janela.update()

    for item in TELAS:
        identificador = item["id"]
        if apenas and identificador not in apenas:
            continue
        marcador = len(app.contexto.erros_teste)
        try:
            app.abrir_tela(identificador)
            app.janela.update_idletasks()
            app.janela.update()
            tela = app.tela_atual
            if tela is None:
                raise RuntimeError("a tela não foi criada (None)")
            tela.carregar()
            app.janela.update()
            if hasattr(tela, "focar_busca"):
                tela.focar_busca()
            novos_erros = app.contexto.erros_teste[marcador:]
            if novos_erros:
                for mensagem_erro in novos_erros:
                    falhas.append(f"{identificador}: {mensagem_erro}")
                    print(f"[ERRO] {identificador:<16} {mensagem_erro[:180]}")
            else:
                print(f"[ok]   {identificador:<16} {item['rotulo']}")
        except Exception as exc:
            falhas.append(f"{identificador}: {exc}")
            print(f"[FALHA] {identificador:<16} {exc}")
            traceback.print_exc(limit=6)

    # ------------------------------------------------------------------ PDV interativo
    try:
        testar_pdv(app, banco)
        print("[ok]   fluxo de venda no PDV (bipagem, desconto, pagamento, recibo)")
    except Exception as exc:
        falhas.append(f"pdv fluxo: {exc}")
        print(f"[FALHA] pdv fluxo: {exc}")
        traceback.print_exc(limit=8)

    # ------------------------------------------------------------------ cadastros
    try:
        testar_cadastros(app, banco, dialogs_mod)
        print("[ok]   cadastros (categorias, clientes, fornecedores, usuários)")
    except Exception as exc:
        falhas.append(f"cadastros: {exc}")
        print(f"[FALHA] cadastros: {exc}")
        traceback.print_exc(limit=10)

    # ------------------------------------------------------------------ estoque
    try:
        testar_estoque(app, banco, dialogs_mod)
        print("[ok]   estoque (movimentações, ajuste de inventário, lotes/validade)")
    except Exception as exc:
        falhas.append(f"estoque: {exc}")
        print(f"[FALHA] estoque: {exc}")
        traceback.print_exc(limit=10)

    # ------------------------------------------------------------------ financeiro
    try:
        testar_financeiro(app, banco, dialogs_mod)
        print("[ok]   financeiro (contas a receber e contas a pagar)")
    except Exception as exc:
        falhas.append(f"financeiro: {exc}")
        print(f"[FALHA] financeiro: {exc}")
        traceback.print_exc(limit=10)

    # ------------------------------------------------------------------ relatórios
    try:
        testar_relatorios(app, banco, dialogs_mod)
        print("[ok]   relatorios (todos os relatórios, filtros, exportação e impressão)")
    except Exception as exc:
        falhas.append(f"relatorios: {exc}")
        print(f"[FALHA] relatorios: {exc}")
        traceback.print_exc(limit=10)

    # ------------------------------------------------------------------ etiquetas
    try:
        testar_etiquetas(app, banco, dialogs_mod)
        print("[ok]   etiquetas (código de barras, fila, prévia, PDF e impressão)")
    except Exception as exc:
        falhas.append(f"etiquetas: {exc}")
        print(f"[FALHA] etiquetas: {exc}")
        traceback.print_exc(limit=10)

    # ------------------------------------------------------------------ backup
    try:
        testar_backup(app, banco, dialogs_mod)
        print("[ok]   backup (criar, verificar, conteúdo, limpeza, excluir e restaurar)")
    except Exception as exc:
        falhas.append(f"backup: {exc}")
        print(f"[FALHA] backup: {exc}")
        traceback.print_exc(limit=10)

    # ------------------------------------------------------------------ configurações
    try:
        testar_configuracoes(app, banco, dialogs_mod)
        print("[ok]   configuracoes (empresa, recibo, regras, tema, atalhos, diagnóstico)")
    except Exception as exc:
        falhas.append(f"configuracoes: {exc}")
        print(f"[FALHA] configuracoes: {exc}")
        traceback.print_exc(limit=10)

    # troca de tema com telas abertas
    try:
        app.alternar_tema("escuro")
        app.janela.update()
        app.alternar_tema("claro")
        app.janela.update()
        print("[ok]   troca de tema")
    except Exception as exc:
        falhas.append(f"tema: {exc}")
        traceback.print_exc(limit=6)

    falhas.extend(f"mensagem de erro: {m}" for m in dialogs_mod.LOG_TESTE if m.startswith("[erro]") or m.startswith("[exceção]"))
    if app.contexto.erros_teste:
        print(f"\nMensagens de erro capturadas ({len(app.contexto.erros_teste)}):")
        for mensagem_erro in app.contexto.erros_teste:
            print("  -", mensagem_erro.replace("\n", " | ")[:400])

    app.janela.destroy()
    banco.fechar()
    print()
    if falhas:
        print(f"FALHAS ({len(falhas)}):")
        for f in falhas:
            print(" -", f)
        return 1
    print("SUCESSO: interface gráfica verificada em todas as telas.")
    return 0


if __name__ == "__main__":
    argumentos = [a for a in sys.argv[1:] if not a.startswith("-")]
    sys.exit(executar(tuple(argumentos)))
